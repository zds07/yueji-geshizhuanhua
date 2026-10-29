"""
parse_brief.py — 从甲方原稿（docx / xlsx）抽出结构化 dict

金标准 xlsx 列位置（openpyxl 1-based）：
    col 2 (B): 镜号 / 幕标题
    col 4 (D): 拍摄地点
    col 5 (E): 画面内容
    col 7 (G): 镜头设计
    col 8 (H): 台词/字幕
    col 12 (L): 时长
    col 13 (M): 备注
    col 14 (N): 道具
    col 15 (O): 音效
    col 16 (P): 参考画面
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path


# xlsx 列位置常量
X_COL_SHOT_NO = 2
X_COL_LOCATION = 4
X_COL_ACTION = 5
X_COL_CAMERA = 7
X_COL_DIALOGUE = 8
X_COL_DURATION = 12
X_COL_NOTES = 13
X_COL_PROPS = 14
X_COL_SFX = 15
X_COL_REFERENCE = 16


def parse_xlsx(path: str) -> dict:
    """从 xlsx 读取制作版或简化版"""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    all_rows = []
    for row in ws.iter_rows(values_only=True):
        all_rows.append(["" if c is None else str(c) for c in row])

    if not all_rows:
        return {"meta": {}, "acts": []}

    # 找信息头起始行
    meta = {"video_name": "", "video_topic": "", "props": "", "product": ""}
    title_row_idx = None
    for i, r in enumerate(all_rows[:15]):
        row_str = " ".join(c for c in r)
        if "短视频拍摄脚本" in row_str:
            title_row_idx = i
            break
        if i == 0 and r and r[0] and "脚本" in r[0]:
            title_row_idx = i
            break

    if title_row_idx is not None:
        # 制作版：信息头在 title_row_idx+1 到 title_row_idx+7
        for offset in range(1, 8):
            if title_row_idx + offset >= len(all_rows):
                break
            r = all_rows[title_row_idx + offset]
            row_str = " ".join(c for c in r)
            for label, key in [("视频名称", "video_name"), ("视频主题", "video_topic"),
                                ("拍摄道具", "props"), ("主要产品", "product")]:
                if label in row_str:
                    # 找标签列后的第一个非空值
                    label_idx = -1
                    for j, c in enumerate(r):
                        if label in c:
                            label_idx = j
                            break
                    if label_idx >= 0:
                        for j in range(label_idx + 1, len(r)):
                            if r[j].strip():
                                meta[key] = r[j].strip()
                                break

    # 找表头行（含"镜号"的行）
    header_row_idx = None
    for i, r in enumerate(all_rows):
        if any("镜号" in str(c) for c in r):
            header_row_idx = i
            break

    if header_row_idx is None:
        # 没表头 → 简化版 xlsx（第一行是标题，第二行是表头）
        if len(all_rows) >= 2:
            meta["video_name"] = all_rows[0][0] if all_rows[0] else meta["video_name"]
            header_row_idx = 1

    # 找镜头起始行
    acts = []
    current_act = None
    shot_no = 0

    for r in all_rows[header_row_idx + 1:]:
        if not r:
            continue
        # 第一列判断：镜号/幕标题
        first = (r[X_COL_SHOT_NO - 1] if len(r) > X_COL_SHOT_NO - 1 else "").strip() if X_COL_SHOT_NO - 1 < len(r) else ""

        # 幕标题行（含"幕"但不是数字）
        if first and not re.match(r"^0?\d+(\.\d+)?$", first) and (
            "幕" in first or first.startswith("【") or
            first.startswith("第一") or first.startswith("第二") or
            first.startswith("第三") or first.startswith("篇章")
        ):
            if len(first) < 30:  # 幕标题不会很长
                current_act = {"name": first, "shots": []}
                acts.append(current_act)
                continue

        # 镜号
        if re.match(r"^0?\d+(\.\d+)?$", first):
            try:
                shot_no = int(float(first))
            except ValueError:
                shot_no += 1
        else:
            continue

        if current_act is None:
            current_act = {"name": None, "shots": []}
            acts.append(current_act)

        shot = {
            "shot_no": shot_no,
            "location": _cell(r, X_COL_LOCATION),
            "action": _cell(r, X_COL_ACTION),
            "camera": _cell(r, X_COL_CAMERA),
            "dialogue": _cell(r, X_COL_DIALOGUE),
            "duration": _cell(r, X_COL_DURATION),
            "notes": _cell(r, X_COL_NOTES),
            "props": _cell(r, X_COL_PROPS),
            "sfx": _cell(r, X_COL_SFX),
            "reference": _cell(r, X_COL_REFERENCE),
        }
        current_act["shots"].append(shot)

    return {"meta": meta, "acts": acts}


def _cell(row, idx):
    """安全获取 row[idx-1] 的值"""
    if idx - 1 < len(row):
        return row[idx - 1].strip() if row[idx - 1] else ""
    return ""


def parse_docx_table(path: str) -> dict:
    """从 docx 表格抽出结构"""
    from docx import Document

    doc = Document(path)
    if not doc.tables:
        return parse_docx_text(path)

    table = doc.tables[0]
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        rows.append(cells)

    if len(rows) < 2:
        return {"meta": {}, "acts": []}

    meta = {"video_name": "", "video_topic": ""}
    if rows and rows[0] and rows[0][0]:
        title = rows[0][0]
        if title.startswith("《"):
            meta["video_name"] = title.strip("《》")
        else:
            meta["video_name"] = title

    header = rows[1]
    col_map = _map_columns(header)

    acts = []
    current_act = None
    shot_no = 0

    for row in rows[2:]:
        if not any(c for c in row):
            continue
        first = (row[0] if row else "").strip()

        if first and (first.startswith("【") or "幕" in first[:10]):
            current_act = {"name": first, "shots": []}
            acts.append(current_act)
            continue

        if first and re.match(r"^0?\d+$", first):
            shot_no = int(first.lstrip("0") or "0") or (shot_no + 1)
        else:
            shot_no += 1

        if current_act is None:
            current_act = {"name": None, "shots": []}
            acts.append(current_act)

        shot = {
            "shot_no": shot_no,
            "location": "", "action": "", "camera": "", "dialogue": "",
            "duration": "", "notes": "", "props": "", "sfx": "", "reference": "",
        }
        for idx, value in enumerate(row):
            if not value:
                continue
            col_name = col_map.get(idx, "")
            if col_name in shot:
                shot[col_name] = value
        current_act["shots"].append(shot)

    return {"meta": meta, "acts": acts}


def _map_columns(header: list[str]) -> dict[int, str]:
    mapping = {
        "序号": "shot_no", "镜号": "shot_no",
        "拍摄地点": "location", "场景": "location", "地点": "location",
        "画面": "action", "画面内容": "action", "画面描述": "action",
        "分镜画面": "camera", "镜头": "camera", "景别": "camera", "镜头设计": "camera",
        "字幕": "dialogue", "字幕/同期声": "dialogue", "台词": "dialogue",
        "旁白": "dialogue", "台词/字幕": "dialogue", "人物对白/旁白": "dialogue",
        "时间": "duration", "时长": "duration",
        "备注": "notes",
        "音效": "sfx", "背景音/音效": "sfx",
        "道具": "props", "人员及所需道具": "props",
        "参考": "reference", "参考画面": "reference",
    }
    result = {}
    for i, h in enumerate(header):
        h = (h or "").strip()
        if h in mapping:
            result[i] = mapping[h]
    return result


def parse_docx_text(path: str) -> dict:
    """从 docx 自由文本抽出结构"""
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    meta = {"video_name": "", "video_topic": "", "shoot_date": "", "deadline": "",
            "props": "", "product": ""}

    if paragraphs and paragraphs[0].startswith("《"):
        meta["video_name"] = paragraphs[0].strip("《》")
    elif paragraphs and paragraphs[0].startswith("【"):
        meta["video_name"] = paragraphs[0].strip("【】")

    for p in paragraphs[:5]:
        if p.startswith("主题："):
            meta["video_topic"] = p.replace("主题：", "").strip()
        if p.startswith("风格："):
            meta["video_topic"] = (meta["video_topic"] + " | " + p.replace("风格：", "").strip()).strip(" |")

    acts = []
    current_act = None
    current_shot = None
    shot_no = 0

    act_pattern = re.compile(r"【(开场|第一幕|第二幕|第三幕|第四幕|篇章[一二三四五六七八九十]|救赎线[ABC]|深渊线[ABC]|第[一二三四五六七八九十]+幕)】")
    scene_pattern = re.compile(r"^(场景|地点|时间)：(.+)$")
    camera_pattern = re.compile(r"^(镜头|景别|运镜|机位|画面)：(.+)$")
    dialogue_pattern = re.compile(r"^([\u4e00-\u9fa5A-Za-z][\u4e00-\u9fa5A-Za-z0-9·]{0,8})(：|:)(.+)$")

    for p in paragraphs:
        if p.startswith("《") or p.startswith("主题：") or p.startswith("风格：") or p.startswith("时长："):
            continue

        m_act = act_pattern.match(p)
        if m_act:
            current_act = {"name": m_act.group(0), "shots": []}
            acts.append(current_act)
            current_shot = None
            continue

        m_scene = scene_pattern.match(p)
        if m_scene:
            if current_shot is None:
                shot_no += 1
                if current_act is None:
                    current_act = {"name": None, "shots": []}
                    acts.append(current_act)
                current_shot = _new_shot(shot_no, location=m_scene.group(2).strip())
                current_act["shots"].append(current_shot)
            else:
                current_shot["location"] = m_scene.group(2).strip()
            continue

        m_camera = camera_pattern.match(p)
        if m_camera:
            if current_shot is None:
                shot_no += 1
                if current_act is None:
                    current_act = {"name": None, "shots": []}
                    acts.append(current_act)
                current_shot = _new_shot(shot_no, camera=m_camera.group(2).strip())
                current_act["shots"].append(current_shot)
            else:
                current_shot["camera"] = (current_shot["camera"] + " | " + m_camera.group(2).strip()).strip(" |")
            continue

        m_dialogue = dialogue_pattern.match(p)
        if m_dialogue:
            role = m_dialogue.group(1).strip()
            text = m_dialogue.group(3).strip()
            if role in ("镜头", "场景"):
                if current_shot is None:
                    shot_no += 1
                    if current_act is None:
                        current_act = {"name": None, "shots": []}
                        acts.append(current_act)
                    current_shot = _new_shot(shot_no, action=text)
                    current_act["shots"].append(current_shot)
                else:
                    current_shot["action"] = (current_shot["action"] + "\n" + text).strip()
            else:
                dialogue = f"{role}：{text}"
                if current_shot is None:
                    shot_no += 1
                    if current_act is None:
                        current_act = {"name": None, "shots": []}
                        acts.append(current_act)
                    current_shot = _new_shot(shot_no, dialogue=dialogue)
                    current_act["shots"].append(current_shot)
                else:
                    if current_shot["dialogue"]:
                        current_shot["dialogue"] += "\n" + dialogue
                    else:
                        current_shot["dialogue"] = dialogue
            continue

        if current_shot is None:
            shot_no += 1
            if current_act is None:
                current_act = {"name": None, "shots": []}
                acts.append(current_act)
            current_shot = _new_shot(shot_no, action=p)
            current_act["shots"].append(current_shot)
        else:
            current_shot["action"] = (current_shot["action"] + "\n" + p).strip()

    return {"meta": meta, "acts": acts}


def _new_shot(shot_no: int, location: str = "", action: str = "", camera: str = "",
              dialogue: str = "", duration: str = "", notes: str = "",
              props: str = "", sfx: str = "", reference: str = "") -> dict:
    return {
        "shot_no": shot_no, "location": location, "action": action, "camera": camera,
        "dialogue": dialogue, "duration": duration, "notes": notes,
        "props": props, "sfx": sfx, "reference": reference,
    }


def parse_brief(path: str) -> dict:
    """主入口：根据文件类型自动选择解析方式"""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(path)
    ext = p.suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        if doc.tables:
            return parse_docx_table(path)
        return parse_docx_text(path)
    elif ext in (".xlsx", ".xls"):
        return parse_xlsx(path)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")


def _cli():
    ap = argparse.ArgumentParser(description="从甲方稿抽出结构化 dict")
    ap.add_argument("--input", required=True, help="原稿路径（docx/xlsx）")
    ap.add_argument("--output", help="输出 JSON 路径（默认打印到 stdout）")
    args = ap.parse_args()
    data = parse_brief(args.input)
    out = json.dumps(data, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(out, encoding="utf-8")
        print(f"OK: {args.output} ({len(out)} chars)")
    else:
        print(out)


if __name__ == "__main__":
    _cli()
